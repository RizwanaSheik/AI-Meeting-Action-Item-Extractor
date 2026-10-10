import json
from io import BytesIO
from pathlib import Path

import pandas as pd
import streamlit as st
from docx import Document
from google.genai import errors as genai_errors
from pypdf import PdfReader

from src.preprocessing import clean_transcript
from src.pipeline import process_meeting_transcript
from src.evaluation import load_annotations, evaluate_extraction


# ==================================================
# PAGE CONFIGURATION
# ==================================================

st.set_page_config(
    page_title="AI Meeting Action-Item Extractor",
    page_icon="📋",
    layout="wide",
)

st.title("📋 AI Meeting Action-Item Extractor")

st.write(
    "Upload a TXT, PDF, or DOCX meeting transcript to extract "
    "structured action items, validate them, evaluate extraction "
    "performance, and identify tasks that need human review."
)


# ==================================================
# FILE EXTRACTION FUNCTIONS
# ==================================================

def extract_txt_text(file_bytes: bytes) -> str:
    """Read a UTF-8 text file."""
    return file_bytes.decode("utf-8-sig")


def extract_pdf_text(file_bytes: bytes) -> str:
    """Extract selectable text from a PDF using pypdf."""
    reader = PdfReader(BytesIO(file_bytes))

    if reader.is_encrypted:
        try:
            result = reader.decrypt("")
        except Exception as error:
            raise ValueError(
                "This PDF is password-protected. "
                "Please upload an unlocked PDF."
            ) from error

        if not result:
            raise ValueError(
                "This PDF is password-protected. "
                "Please upload an unlocked PDF."
            )

    extracted_pages = []

    for page in reader.pages:
        extracted_pages.append(page.extract_text() or "")

    return "\n".join(extracted_pages)


def extract_docx_text(file_bytes: bytes) -> str:
    """Extract paragraphs and table contents from a DOCX."""
    document = Document(BytesIO(file_bytes))
    extracted_parts = []

    # Extract paragraphs.
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            extracted_parts.append(text)

    # Extract tables.
    for table in document.tables:
        for row in table.rows:
            cells = [
                cell.text.strip().replace("\n", " ")
                for cell in row.cells
            ]

            if any(cells):
                extracted_parts.append(" | ".join(cells))

    return "\n".join(extracted_parts)


def extract_transcript_text(uploaded_file) -> str:
    """Extract text according to the uploaded file extension."""
    file_bytes = uploaded_file.getvalue()
    extension = uploaded_file.name.lower().rsplit(".", 1)[-1]

    if extension == "txt":
        return extract_txt_text(file_bytes)

    if extension == "pdf":
        return extract_pdf_text(file_bytes)

    if extension == "docx":
        return extract_docx_text(file_bytes)

    raise ValueError(
        "Unsupported file type. Please upload TXT, PDF, or DOCX."
    )


# ==================================================
# TRANSCRIPT UPLOAD
# ==================================================

uploaded_file = st.file_uploader(
    "Upload a meeting transcript",
    type=["txt", "pdf", "docx"],
)

if uploaded_file is None:
    st.info("Upload a TXT, PDF, or DOCX transcript to get started.")
    st.stop()


# ==================================================
# READ AND CLEAN TRANSCRIPT
# ==================================================

try:
    raw_transcript = extract_transcript_text(uploaded_file)

except UnicodeDecodeError:
    st.error(
        "The TXT file could not be decoded as UTF-8. "
        "Please save it using UTF-8 encoding and upload it again."
    )
    st.stop()

except ValueError as error:
    st.error(str(error))
    st.stop()

except Exception as error:
    st.error(
        "The uploaded file could not be read. "
        "Check that it is a valid, uncorrupted document."
    )
    st.caption(f"Details: {error}")
    st.stop()


transcript = clean_transcript(raw_transcript)

if not transcript.strip():
    st.error(
        "No extractable text was found in this document. "
        "If this is a scanned PDF, OCR is required before processing."
    )
    st.stop()


with st.expander("View cleaned transcript"):
    st.text(transcript)


# ==================================================
# EXTRACTION AND VALIDATION
# ==================================================

if st.button("Extract Action Items", type="primary"):

    try:
        with st.spinner("Processing transcript with Gemini..."):
            results = process_meeting_transcript(transcript)

    except genai_errors.APIError as error:
        st.error(
            "Gemini could not process the transcript. "
            "The service may be temporarily unavailable, or your "
            "API request may have a configuration or quota issue."
        )
        st.caption(f"API error details: {error}")
        st.stop()

    except (json.JSONDecodeError, ValueError) as error:
        st.error(
            "The extraction result could not be validated. "
            "Please check your extraction code and try again."
        )
        st.caption(f"Details: {error}")
        st.stop()

    except Exception as error:
        st.error(
            "An unexpected error occurred while processing the transcript."
        )
        st.caption(f"Error type: {type(error).__name__}")
        st.caption(f"Details: {error}")
        st.stop()

    # Retrieve results.
    action_items = results.get("action_items", [])
    review_queue = results.get("review_queue", [])

    passed_count = len(action_items) - len(review_queue)

    # ==================================================
    # RESULTS SUMMARY
    # ==================================================

    st.subheader("Results Summary")

    col1, col2, col3 = st.columns(3)

    col1.metric(
        "Items Extracted",
        results.get("total_extracted", len(action_items)),
    )

    col2.metric(
        "Needs Review",
        results.get("total_needing_review", len(review_queue)),
    )

    col3.metric(
        "Passed Initial Validation",
        passed_count,
    )

    # ==================================================
    # VALIDATED ACTION ITEMS
    # ==================================================

    st.subheader("Validated Action Items")

    if action_items:
        df = pd.DataFrame(action_items)

        st.dataframe(
            df,
            width="stretch",
            hide_index=True,
        )

        # Download action items as CSV.
        csv_data = df.to_csv(index=False).encode("utf-8-sig")

        st.download_button(
            label="Download CSV",
            data=csv_data,
            file_name="action_items.csv",
            mime="text/csv",
        )

        # Download action items as JSON.
        json_data = json.dumps(
            action_items,
            indent=2,
            ensure_ascii=False,
        )

        st.download_button(
            label="Download JSON",
            data=json_data,
            file_name="action_items.json",
            mime="application/json",
        )

    else:
        st.info("No action items were found.")

    # ==================================================
    # EVALUATION METRICS
    # ==================================================

    st.subheader("Extraction Evaluation")

    annotation_folder = Path("data/annotations")

    # Match the annotation JSON to the uploaded filename.
    #
    # Examples:
    # meeting_001.txt                  -> meeting_001.json
    # example_meeting_transcript.pdf   -> example_meeting_transcript.json
    # example_meeting_transcript.docx  -> example_meeting_transcript.json

    file_stem = Path(uploaded_file.name).stem
    annotation_path = annotation_folder / f"{file_stem}.json"

    if annotation_path.is_file():

        try:
            reference_items = load_annotations(annotation_path)

            evaluation = evaluate_extraction(
                action_items,
                reference_items,
            )

            metric1, metric2, metric3 = st.columns(3)

            metric1.metric(
                "Precision",
                f"{evaluation['precision'] * 100:.2f}%",
            )

            metric2.metric(
                "Recall",
                f"{evaluation['recall'] * 100:.2f}%",
            )

            metric3.metric(
                "F1 Score",
                f"{evaluation['f1_score']:.4f}",
            )

            st.caption(
                "Evaluation compares extracted action items against "
                "reference annotations using task and owner matching. "
                "Scores depend on the annotation quality and matching threshold."
            )

            with st.expander("View evaluation details"):

                st.write(
                    f"Reference items: {len(reference_items)}"
                )

                st.write(
                    f"True positives: {evaluation['true_positives']}"
                )

                st.write(
                    f"False positives: {evaluation['false_positives']}"
                )

                st.write(
                    f"False negatives: {evaluation['false_negatives']}"
                )

                matched_items = evaluation.get("matched_items", [])

                if matched_items:
                    st.write("Matched Items")

                    st.dataframe(
                        pd.DataFrame(matched_items),
                        width="stretch",
                        hide_index=True,
                    )

        except (OSError, ValueError, KeyError, TypeError) as error:
            st.warning(
                "Evaluation could not be completed. "
                "Check the annotation JSON and evaluation module."
            )
            st.caption(f"Details: {error}")

    else:
        st.info(
            f"No reference annotations found for {uploaded_file.name}. "
            f"Create {file_stem}.json inside data/annotations "
            "to enable evaluation for this transcript."
        )

    # ==================================================
    # HUMAN REVIEW QUEUE
    # ==================================================

    st.subheader("Human Review Queue")

    if review_queue:
        st.warning(
            "These action items have validation issues "
            "and need human review."
        )

        review_df = pd.DataFrame(review_queue)

        st.dataframe(
            review_df,
            width="stretch",
            hide_index=True,
        )

        review_csv = review_df.to_csv(
            index=False
        ).encode("utf-8-sig")

        st.download_button(
            label="Download Review Queue CSV",
            data=review_csv,
            file_name="human_review_queue.csv",
            mime="text/csv",
        )

    else:
        st.success("No items currently require review.")